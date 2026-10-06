import os
from docx import Document

OUTPUT_FOLDER = "output"
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

def save_doc(summary, decisions, actions, agenda, entities, filename):

    doc = Document()
    doc.add_heading("AI Meeting Minutes", 0)

    doc.add_heading("Summary", 1)
    doc.add_paragraph(summary)

    doc.add_heading("Agenda", 1)
    for a in agenda:
        doc.add_paragraph(a, style="List Bullet")

    doc.add_heading("Decisions", 1)
    for d in decisions:
        doc.add_paragraph(d, style="List Bullet")

    doc.add_heading("Action Items", 1)
    for a in actions:
        doc.add_paragraph(a, style="List Bullet")

    doc.add_heading("Entities", 1)
    for e in entities:
        doc.add_paragraph(e)

    path = os.path.join(OUTPUT_FOLDER, filename + "_minutes.docx")
    doc.save(path)

    return path